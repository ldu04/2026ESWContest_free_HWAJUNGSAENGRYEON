"""_deck_edit.py — 덱 텍스트 편집 헬퍼. **서식을 절대 잃지 않는 것**이 유일한 목적.

python-pptx 에서 `shape.text = "..."` 나 `run.text = ...` 를 잘못 쓰면 글꼴·크기·색이
기본값으로 돌아간다. 여기 있는 함수들은 **기존 run 의 rPr(서식) 을 복제해서** 새 run 을 만든다.

줄바꿈은 새 문단이 아니라 `<a:br/>`(연성 줄바꿈)로 넣는다 — 문단을 나누면 문단 간격이
붙어 레이아웃이 밀린다.
"""
from __future__ import annotations
import copy
from pptx.util import Pt

NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _first_run(para):
    if not para.runs:
        para.add_run()
    return para.runs[0]


def set_lines(para, lines):
    """문단 내용을 lines 로 교체하고 줄 사이에 <a:br/> 를 넣는다. 첫 run 서식을 물려준다."""
    if isinstance(lines, str):
        raise TypeError("set_lines 는 리스트를 받는다. 문자열을 주면 글자마다 줄바꿈이 된다 "
                        "— 실제로 그 사고가 났다(2026-09-02, s13·s21). set_para() 를 쓸 것.")
    r0 = _first_run(para)
    proto = copy.deepcopy(r0._r)
    for r in para.runs[1:]:
        r._r.getparent().remove(r._r)
    for br in para._p.findall(NS + "br"):
        para._p.remove(br)
    r0.text = lines[0]
    anchor = r0._r
    for ln in lines[1:]:
        br = anchor.makeelement(NS + "br", {})
        anchor.addnext(br)
        new = copy.deepcopy(proto)
        for t in new.findall(NS + "t"):
            t.text = ln
        br.addnext(new)
        anchor = new


def break_before(para, needle):
    """문단 텍스트에서 needle 앞을 잘라 두 줄로 만든다. 못 찾으면 False."""
    txt = para.text
    i = txt.find(needle)
    if i <= 0:
        return False
    set_lines(para, [txt[:i].rstrip(), txt[i:]])
    return True


def replace_in_para(para, old, new):
    """문단 안의 문자열 치환. run 경계를 넘어도 되도록 문단 단위로 다시 쓴다."""
    txt = para.text
    if old not in txt:
        return False
    parts = txt.replace(old, new).split("\n")
    set_lines(para, parts)
    return True


def set_para(para, text):
    set_lines(para, [text])


def insert_para_after(tf, i, text, size_pt=None):
    """i 번 문단을 복제해 뒤에 끼운다 — 글머리·들여쓰기·글꼴이 그대로 따라온다."""
    from pptx.text.text import _Paragraph
    src = tf.paragraphs[i]._p
    new = copy.deepcopy(src)
    src.addnext(new)
    p = _Paragraph(new, tf)
    for br in p._p.findall(NS + "br"):
        p._p.remove(br)
    set_lines(p, [text])
    if size_pt:
        for r in p.runs:
            r.font.size = Pt(size_pt)
    return p


def cell_para(tbl, r, c, i=0):
    return tbl.cell(r, c).text_frame.paragraphs[i]


def set_bullets(tf, lines):
    """글머리 목록을 **문단 하나씩**으로 만든다.

    왜 <a:br/> 를 쓰면 안 되나(2026-09-02 사고): 이 덱의 본문 pPr 은 내어쓰기
    (marL>0 · indent<0)다. 문단의 **첫 줄**만 marL+indent 에서 시작하고 나머지 줄은
    marL 에서 시작한다. <a:br/> 로 이은 둘째 글머리는 「나머지 줄」로 취급돼
    **첫 글머리보다 안쪽으로 밀린다.** 문단을 나누면 각 글머리가 첫 줄이 된다.
    """
    import copy as _c
    keep = tf.paragraphs[0]
    proto = _c.deepcopy(keep._p)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for br in keep._p.findall(NS + "br"):
        keep._p.remove(br)
    set_lines(keep, [lines[0]])
    anchor = keep._p
    from pptx.text.text import _Paragraph
    for ln in lines[1:]:
        new = _c.deepcopy(proto)
        anchor.addnext(new)
        p = _Paragraph(new, tf)
        for br in p._p.findall(NS + "br"):
            p._p.remove(br)
        set_lines(p, [ln])
        anchor = new
